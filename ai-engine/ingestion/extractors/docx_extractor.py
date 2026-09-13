"""DOCX Extractor using python-docx for Word document processing."""

import io
import logging
from typing import List
import docx
from schemas.document import ExtractedDocument, Section
from schemas.base import generate_id
from .base import BaseExtractor

logger = logging.getLogger("arcana.ingestion.docx")


class DocxExtractor(BaseExtractor):
    """Deterministic extractor for DOCX documents."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        sections: List[Section] = []
        current_title = "Overview"
        current_level = 1
        current_paragraphs: List[str] = []
        all_text_parts: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            all_text_parts.append(text)
            style_name = p.style.name if p.style else ""

            if "Heading" in style_name:
                if current_paragraphs:
                    sections.append(
                        Section(
                            section_id=generate_id("sec"),
                            title=current_title,
                            level=current_level,
                            content="\n".join(current_paragraphs).strip(),
                        )
                    )
                    current_paragraphs = []
                try:
                    current_level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    current_level = 1
                current_title = text
            else:
                current_paragraphs.append(text)

        # Extract tables as structured text
        for table in doc.tables:
            table_lines: List[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(cells))
            if table_lines:
                table_text = "\n".join(table_lines)
                current_paragraphs.append(f"\n[Table]\n{table_text}\n")
                all_text_parts.append(table_text)

        if current_paragraphs:
            sections.append(
                Section(
                    section_id=generate_id("sec"),
                    title=current_title,
                    level=current_level,
                    content="\n".join(current_paragraphs).strip(),
                )
            )

        full_text = "\n\n".join(all_text_parts)
        title = filename.rsplit(".", 1)[0].replace("_", " ").title()

        return ExtractedDocument(
            document_id=generate_id("doc"),
            filename=filename,
            source_type="docx",
            title=title,
            sections=sections,
            raw_text=full_text,
            metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
        )
