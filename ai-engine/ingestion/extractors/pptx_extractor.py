"""PPTX Extractor using python-pptx for PowerPoint presentation processing."""

import io
import logging
from typing import List
import pptx
from schemas.document import ExtractedDocument, Section
from schemas.base import generate_id
from .base import BaseExtractor

logger = logging.getLogger("arcana.ingestion.pptx")


class PPTXExtractor(BaseExtractor):
    """Deterministic extractor for PPTX presentation decks."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        stream = io.BytesIO(file_bytes)
        prs = pptx.Presentation(stream)

        sections: List[Section] = []
        all_text_parts: List[str] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_title = f"Slide {slide_num}"
            slide_texts: List[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

            # Check if slide has a dedicated title shape
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()

            if slide_texts:
                combined_text = "\n".join(slide_texts)
                all_text_parts.append(combined_text)
                sections.append(
                    Section(
                        section_id=generate_id("sec"),
                        title=slide_title,
                        level=1,
                        content=combined_text,
                        page_start=slide_num,
                        page_end=slide_num,
                        metadata={"slide_number": slide_num},
                    )
                )

        full_text = "\n\n".join(all_text_parts)
        title = filename.rsplit(".", 1)[0].replace("_", " ").title()

        return ExtractedDocument(
            document_id=generate_id("doc"),
            filename=filename,
            source_type="pptx",
            title=title,
            sections=sections,
            raw_text=full_text,
            metadata={"slide_count": len(prs.slides)},
        )
