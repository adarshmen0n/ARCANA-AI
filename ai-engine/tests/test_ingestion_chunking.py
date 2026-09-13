"""Unit tests for ingestion, extraction, preprocessing, and semantic chunking."""

import io
import pytest
from pypdf import PdfWriter
import docx
from ingestion.validator import DocumentValidator, DocumentValidationError
from ingestion.engine import DocumentIngestionEngine
from preprocessing.cleaner import TextCleaner
from chunking.semantic_chunker import SemanticChunker
from schemas.document import ExtractedDocument, Section


def test_validator_rejects_empty():
    with pytest.raises(DocumentValidationError, match="is empty"):
        DocumentValidator.validate(b"", "lesson.txt")


def test_validator_rejects_unsupported_extension():
    with pytest.raises(DocumentValidationError, match="Unsupported format"):
        DocumentValidator.validate(b"content", "lesson.exe")


def test_validator_rejects_fake_pdf():
    with pytest.raises(DocumentValidationError, match="valid PDF header"):
        DocumentValidator.validate(b"not a real pdf", "lesson.pdf")


def test_txt_ingestion_and_sections():
    engine = DocumentIngestionEngine()
    txt_content = (
        "# Introduction to Python\n"
        "Python is a high-level programming language.\n\n"
        "## Variables\n"
        "Variables store data values in memory.\n"
        "x = 5\n"
    ).encode("utf-8")

    doc = engine.ingest(txt_content, "python_intro.txt")
    assert doc.source_type == "txt"
    assert len(doc.sections) == 2
    assert doc.sections[0].title == "Introduction to Python"
    assert doc.sections[1].title == "Variables"
    assert "x = 5" in doc.sections[1].content


def test_pdf_extraction_page_tracking():
    # Generate a minimal in-memory PDF with 2 pages
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    pdf_buffer = io.BytesIO()
    writer.write(pdf_buffer)
    pdf_bytes = pdf_buffer.getvalue()

    engine = DocumentIngestionEngine()
    doc = engine.ingest(pdf_bytes, "test_sample.pdf")
    assert doc.source_type == "pdf"
    assert doc.metadata["total_pages"] == 2


def test_docx_extraction():
    # Generate a minimal in-memory DOCX
    docx_file = docx.Document()
    docx_file.add_heading("Chapter 1: Operating Systems", level=1)
    docx_file.add_paragraph("An operating system manages computer hardware.")
    docx_file.add_heading("Process Management", level=2)
    docx_file.add_paragraph("A process is a program in execution.")

    buffer = io.BytesIO()
    docx_file.save(buffer)
    docx_bytes = buffer.getvalue()

    engine = DocumentIngestionEngine()
    doc = engine.ingest(docx_bytes, "os_notes.docx")
    assert doc.source_type == "docx"
    assert len(doc.sections) >= 2
    assert "operating system" in doc.raw_text.lower()


def test_text_cleaner_preserves_code_and_removes_noise():
    dirty = (
        "Page 14\n"
        "CPU schedul-\ning is essential.\n\n"
        "```python\n"
        "def run_cpu():\n"
        "    return True\n"
        "```\n\n"
        "15\n"
    )
    cleaned = TextCleaner.clean_text(dirty)
    assert "scheduling is essential." in cleaned
    assert "Page 14" not in cleaned
    assert "def run_cpu():\n    return True" in cleaned


def test_semantic_chunker():
    chunker = SemanticChunker(max_chunk_size=300, min_chunk_size=100, overlap=50)
    sample_doc = ExtractedDocument(
        filename="biology.txt",
        source_type="txt",
        title="Cell Biology",
        sections=[
            Section(
                title="Mitochondria",
                level=1,
                content="Mitochondria are membrane-bound cell organelles. " * 5,
                page_start=1,
                page_end=1,
            )
        ],
    )
    chunks = chunker.chunk_document(sample_doc)
    assert len(chunks) >= 1
    assert chunks[0].chapter == "Cell Biology"
    assert chunks[0].section == "Mitochondria"
    assert chunks[0].chunking_version == "1.0"
    assert chunks[0].page_start == 1
