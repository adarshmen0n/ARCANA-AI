"""DOCX Extractor using python-docx."""

import io
from typing import List
import docx
from app.ingestion.extractors.base import BaseExtractor
from app.ingestion.models import DocumentMetadata, ExtractedDocument, ExtractedPage, DocumentStatus


class DOCXExtractor(BaseExtractor):
    """Extracts text paragraphs, headings, and tables from DOCX files."""

    def extract(self, content: bytes, metadata: DocumentMetadata) -> ExtractedDocument:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)

        text_lines: List[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_lines.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_lines.append(" | ".join(row_cells))

        full_text = "\n\n".join(text_lines)

        # Treat DOCX as a continuous page or estimate sections
        pages = [
            ExtractedPage(
                page_number=1,
                text=full_text,
                character_count=len(full_text),
            )
        ]

        metadata.total_pages = 1
        metadata.total_characters = len(full_text)
        metadata.status = DocumentStatus.EXTRACTED

        return ExtractedDocument(metadata=metadata, full_text=full_text, pages=pages)
