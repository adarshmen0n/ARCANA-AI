"""PPTX Extractor using python-pptx."""

import io
from typing import List
from pptx import Presentation
from app.ingestion.extractors.base import BaseExtractor
from app.ingestion.models import DocumentMetadata, ExtractedDocument, ExtractedPage, DocumentStatus


class PPTXExtractor(BaseExtractor):
    """Extracts text slide-by-slide from PowerPoint presentations."""

    def extract(self, content: bytes, metadata: DocumentMetadata) -> ExtractedDocument:
        stream = io.BytesIO(content)
        prs = Presentation(stream)

        pages: List[ExtractedPage] = []
        all_text_parts: List[str] = []

        for idx, slide in enumerate(prs.slides, start=1):
            slide_lines: List[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            slide_lines.append(" | ".join(row_cells))

            slide_text = "\n".join(slide_lines).strip()
            pages.append(
                ExtractedPage(
                    page_number=idx,
                    text=slide_text,
                    character_count=len(slide_text),
                )
            )
            if slide_text:
                all_text_parts.append(f"--- Slide {idx} ---\n{slide_text}")

        full_text = "\n\n".join(all_text_parts)
        metadata.total_pages = len(pages)
        metadata.total_characters = len(full_text)
        metadata.status = DocumentStatus.EXTRACTED

        return ExtractedDocument(metadata=metadata, full_text=full_text, pages=pages)
