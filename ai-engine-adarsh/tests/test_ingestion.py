"""Tests for Document Ingestion and Extraction Subsystem."""

import io
import os
import sys
import pytest
from fastapi.testclient import TestClient
import docx
from pypdf import PdfWriter

# Ensure ai-engine-adarsh is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.main import app
from app.core.errors import ValidationException
from app.ingestion.models import DocumentStatus, SupportedFormat
from app.ingestion.validators import validate_file_metadata
from app.ingestion.extractors.factory import get_extractor
from app.ingestion.service import ingestion_service

client = TestClient(app)


def test_validator_supported_formats():
    """Verify that supported extensions are accepted and mapped correctly."""
    assert validate_file_metadata("lecture.pdf", 1024)[1] == SupportedFormat.PDF
    assert validate_file_metadata("syllabus.docx", 2048)[1] == SupportedFormat.DOCX
    assert validate_file_metadata("slides.pptx", 4096)[1] == SupportedFormat.PPTX
    assert validate_file_metadata("notes.txt", 512)[1] == SupportedFormat.TXT


def test_validator_unsupported_formats():
    """Verify that unsupported extensions are rejected."""
    with pytest.raises(ValidationException) as exc_info:
        validate_file_metadata("script.exe", 1024)
    assert "Unsupported file format '.exe'" in str(exc_info.value)


def test_validator_empty_file():
    """Verify that empty files are rejected."""
    with pytest.raises(ValidationException) as exc_info:
        validate_file_metadata("empty.pdf", 0)
    assert "is empty" in str(exc_info.value)


def test_validator_oversized_file():
    """Verify that files exceeding 25MB are rejected."""
    oversized = 26 * 1024 * 1024
    with pytest.raises(ValidationException) as exc_info:
        validate_file_metadata("large.pdf", oversized)
    assert "File size exceeds maximum permitted limit" in str(exc_info.value)


def test_txt_extractor():
    """Verify plain text extraction."""
    extractor = get_extractor(SupportedFormat.TXT)
    meta = ingestion_service.register_upload("test.txt", b"CPU Scheduling - FCFS and SJF")
    extracted = extractor.extract(b"CPU Scheduling - FCFS and SJF", meta)

    assert extracted.full_text == "CPU Scheduling - FCFS and SJF"
    assert extracted.metadata.total_characters == len("CPU Scheduling - FCFS and SJF")
    assert extracted.metadata.total_pages == 1
    assert extracted.metadata.status == DocumentStatus.EXTRACTED


def test_docx_extractor():
    """Verify DOCX extraction from an in-memory document."""
    doc = docx.Document()
    doc.add_heading("Operating Systems: CPU Scheduling", level=1)
    doc.add_paragraph("First-Come, First-Served is non-preemptive.")

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    extractor = get_extractor(SupportedFormat.DOCX)
    meta = ingestion_service.register_upload("os_notes.docx", docx_bytes)
    extracted = extractor.extract(docx_bytes, meta)

    assert "Operating Systems: CPU Scheduling" in extracted.full_text
    assert "First-Come, First-Served is non-preemptive." in extracted.full_text
    assert extracted.metadata.status == DocumentStatus.EXTRACTED


def test_pdf_extractor():
    """Verify PDF extraction from an in-memory generated PDF."""
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    extractor = get_extractor(SupportedFormat.PDF)
    meta = ingestion_service.register_upload("sample.pdf", pdf_bytes)
    extracted = extractor.extract(pdf_bytes, meta)

    assert extracted.metadata.total_pages == 1
    assert extracted.metadata.status == DocumentStatus.EXTRACTED


def test_api_document_upload_and_status():
    """Verify full HTTP flow: upload document -> poll status -> fetch content."""
    sample_text = b"Topic: CPU Scheduling\nFCFS Algorithm: executes in order of arrival."
    files = {"file": ("cpu_scheduling.txt", sample_text, "text/plain")}

    # 1. Upload Document
    upload_response = client.post("/api/v1/documents/upload", files=files)
    assert upload_response.status_code == 202
    upload_data = upload_response.json()
    document_id = upload_data["document_id"]
    assert upload_data["status"] == "pending"

    # 2. Check Status (FastAPI TestClient executes BackgroundTasks synchronously)
    status_response = client.get(f"/api/v1/documents/{document_id}/status")
    assert status_response.status_code == 200
    status_data = status_response.json()
    assert status_data["document_id"] == document_id
    assert status_data["status"] == "extracted"
    assert status_data["total_characters"] > 0

    # 3. Retrieve Content
    content_response = client.get(f"/api/v1/documents/{document_id}/content")
    assert content_response.status_code == 200
    content_data = content_response.json()
    assert "CPU Scheduling" in content_data["full_text"]

    # 4. List Documents
    list_response = client.get("/api/v1/documents")
    assert list_response.status_code == 200
    doc_list = list_response.json()
    assert any(d["document_id"] == document_id for d in doc_list)
